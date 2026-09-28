"""Send training reminders from the adjacent Excel workbook using SMTP."""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from datetime import date, datetime, time
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

import smtp_settings


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("APP_DATA_DIR", str(BASE_DIR))).resolve()
WORKBOOK_PATH = DATA_DIR / "training_reminders.xlsx"
LOG_PATH = DATA_DIR / "reminder_automation.log"
TIME_ZONE = ZoneInfo("Asia/Kolkata")
REQUIRED_HEADERS = {
    "recipient email",
    "participant name",
    "reminder date",
    "reminder time",
    "actual training date",
    "status",
    "sent at",
    "gmail message id",
    "notes",
}
SAMPLE_EMAILS = {"replace-me@example.com", "example@example.com"}


def log_setup() -> None:
    logging.basicConfig(
        filename=LOG_PATH,
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
    )


def report(message: str, error: bool = False) -> None:
    stream = sys.stderr if error else sys.stdout
    if stream is None:
        (logging.error if error else logging.info)(message)
        return
    print(message, file=stream)


def normalized(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def find_header_row(sheet) -> tuple[int, dict[str, int]]:
    for row_index in range(1, min(sheet.max_row, 20) + 1):
        headers = {normalized(cell.value): cell.column for cell in sheet[row_index] if cell.value}
        if REQUIRED_HEADERS.issubset(headers):
            return row_index, headers
    raise ValueError(
        "Could not find the required column headers. Use the provided training_reminders.xlsx template."
    )


def as_date(value: Any, label: str, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(value.strip(), pattern).date()
            except ValueError:
                continue
    raise ValueError(f"Row {row_number}: {label} must be an Excel date (or YYYY-MM-DD text).")


def as_time(value: Any, row_number: int) -> time:
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        for pattern in ("%H:%M", "%I:%M %p", "%H:%M:%S", "%I.%M%p", "%I.%M %p"):
            try:
                return datetime.strptime(value.strip(), pattern).time()
            except ValueError:
                continue
    raise ValueError(f"Row {row_number}: Reminder Time must be an Excel time (for example 09:30).")


def due_at(row, headers: dict[str, int], row_number: int) -> datetime:
    reminder_date = as_date(row[headers["reminder date"] - 1], "Reminder Date", row_number)
    reminder_time = as_time(row[headers["reminder time"] - 1], row_number)
    return datetime.combine(reminder_date, reminder_time, tzinfo=TIME_ZONE)


def row_value(row, headers: dict[str, int], name: str) -> Any:
    return row[headers[name] - 1]


def build_message(
    recipient: str,
    participant: str,
    training_date: date,
    topic: str = "",
    department: str = "",
    training_time: str = "",
    trainer: str = "",
) -> EmailMessage:
    message = EmailMessage()
    message["To"] = recipient
    message["Subject"] = f"Training reminder: {topic or 'scheduled session'} - {training_date.strftime('%d %B %Y')}"
    greeting = f"Hello {participant}," if participant else "Hello,"
    details = [
        f"Training topic: {topic}" if topic else "",
        f"Department: {department}" if department else "",
        f"Training date: {training_date.strftime('%A, %d %B %Y')}",
        f"Training time: {training_time}" if training_time else "",
        f"Trainer: {trainer}" if trainer else "",
    ]
    message.set_content(
        f"{greeting}\n\n"
        "This is a reminder about your scheduled training session.\n\n"
        + "\n".join(detail for detail in details if detail)
        + "\n\n"
        "Please make the necessary arrangements to attend.\n\n"
        "Regards,\nTraining Team"
    )
    return message


def save_workbook(workbook, workbook_path: Path) -> None:
    workbook.save(workbook_path)


def process(send: bool, workbook_path: Path) -> int:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    header_row, headers = find_header_row(sheet)
    now = datetime.now(TIME_ZONE)
    due_count = 0

    for row_number in range(header_row + 1, sheet.max_row + 1):
        row = [sheet.cell(row_number, col).value for col in range(1, sheet.max_column + 1)]
        status = str(row_value(row, headers, "status") or "").strip()
        if status.lower() != "ready":
            continue

        recipient = str(row_value(row, headers, "recipient email") or "").strip()
        if send and (recipient.lower() in SAMPLE_EMAILS or not recipient):
            report(f"Row {row_number}: skipped; add a real recipient email before setting Status to Ready.")
            continue
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", recipient):
            sheet.cell(row_number, headers["status"]).value = "Error"
            sheet.cell(row_number, headers["notes"]).value = "Invalid recipient email address."
            logging.error("Row %s: invalid recipient email", row_number)
            continue

        try:
            scheduled = due_at(row, headers, row_number)
            training_date = as_date(
                row_value(row, headers, "actual training date"), "Actual Training Date", row_number
            )
            participant = str(row_value(row, headers, "participant name") or "").strip()
            topic = str(row_value(row, headers, "training topic") or "").strip() if "training topic" in headers else ""
            department = str(row_value(row, headers, "department") or "").strip() if "department" in headers else ""
            training_time = str(row_value(row, headers, "training time") or "").strip() if "training time" in headers else ""
            trainer = str(row_value(row, headers, "trainer name") or "").strip() if "trainer name" in headers else ""
        except ValueError as error:
            sheet.cell(row_number, headers["status"]).value = "Error"
            sheet.cell(row_number, headers["notes"]).value = str(error)
            logging.error("%s", error)
            continue

        if scheduled > now:
            continue

        due_count += 1
        if not send:
            report(
                f"DUE row {row_number}: {recipient} | training {training_date:%Y-%m-%d} "
                f"| reminder was {scheduled:%Y-%m-%d %H:%M %Z}"
            )
            continue

        # Persist an in-progress marker before sending so a crash does not silently resend.
        sheet.cell(row_number, headers["status"]).value = "Sending"
        sheet.cell(row_number, headers["notes"]).value = ""
        save_workbook(workbook, workbook_path)
        try:
            message = build_message(recipient, participant, training_date, topic, department, training_time, trainer)
            smtp_settings.send_message(message)
            sheet.cell(row_number, headers["status"]).value = "Sent"
            sheet.cell(row_number, headers["sent at"]).value = datetime.now(TIME_ZONE).replace(tzinfo=None)
            sheet.cell(row_number, headers["gmail message id"]).value = ""
            sheet.cell(row_number, headers["notes"]).value = ""
            save_workbook(workbook, workbook_path)
            logging.info("Sent reminder for row %s to %s", row_number, recipient)
            report(f"SENT row {row_number}: {recipient}")
        except Exception as error:  # SMTP errors must be recorded for the user.
            sheet.cell(row_number, headers["status"]).value = "Error"
            sheet.cell(row_number, headers["notes"]).value = str(error)[:300]
            save_workbook(workbook, workbook_path)
            logging.exception("Failed to send row %s", row_number)
            report(f"ERROR row {row_number}: {error}", error=True)

    if send:
        # Also persist validation errors when no sendable row made it to the send path.
        save_workbook(workbook, workbook_path)
    if not send:
        if due_count:
            report(f"Dry run only: {due_count} due reminder(s); no email was sent.")
        else:
            report("Dry run complete: no Ready reminders are due. No email was sent.")
    workbook.close()
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Send due training reminders from the Excel workbook.")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true", help="Show due reminders without sending email.")
    mode.add_argument("--send", action="store_true", help="Send due reminders through the configured SMTP account.")
    mode.add_argument("--check-smtp", action="store_true", help="Check SMTP settings without sending email.")
    parser.add_argument(
        "--workbook",
        type=Path,
        default=WORKBOOK_PATH,
        help="Workbook path (defaults to training_reminders.xlsx beside this script).",
    )
    args = parser.parse_args()
    log_setup()
    try:
        if args.check_smtp:
            report(smtp_settings.check_connection())
            return 0
        return process(send=args.send, workbook_path=args.workbook.resolve())
    except Exception as error:
        logging.exception("Automation stopped")
        report(f"Automation stopped: {error}", error=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
